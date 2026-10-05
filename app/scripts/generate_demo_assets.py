"""Genera i documenti demo fittizi (§11 della spec) in app/demo_assets/.

- busta_paga_demo.pdf      -> entrata principale, gergo: INPS, IRPEF, TFR...
- estratto_conto_demo.pdf  -> uscite mensili, gergo: TAEG, interessi di mora...
- foglio_spese_demo.xlsx   -> voci annuali da normalizzare (RF-02)

Tutti i dati vengono da finsup/demo_data.py (persona e aziende inventate).
Uso, dalla root del repo:  .venv/Scripts/python app/scripts/generate_demo_assets.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font

APP_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_DIR))

from finsup import demo_data as d  # noqa: E402
from finsup.budget import eur  # noqa: E402

OUT_DIR = APP_DIR / "demo_assets"
WIN_FONT = Path("C:/Windows/Fonts/arial.ttf")  # serve un TTF per il simbolo €


def _pdf() -> FPDF:
    pdf = FPDF()
    pdf.add_page()
    if WIN_FONT.exists():
        pdf.add_font("Body", fname=str(WIN_FONT))
        pdf.add_font("Body", style="B", fname=str(WIN_FONT.with_name("arialbd.ttf")))
        pdf.font_name_body = "Body"
    else:
        pdf.font_name_body = "Helvetica"
    return pdf


def _text(pdf: FPDF, text: str, size: int = 10, bold: bool = False, h: float = 6) -> None:
    if pdf.font_name_body == "Helvetica":
        text = text.replace("€", "EUR")
    pdf.set_font(pdf.font_name_body, "B" if bold else "", size)
    pdf.multi_cell(0, h, text, new_x="LMARGIN", new_y="NEXT")


def _table(pdf: FPDF, rows: list[tuple[str, ...]], widths: tuple[int, ...]) -> None:
    for i, row in enumerate(rows):
        pdf.set_font(pdf.font_name_body, "B" if i == 0 else "", 9)
        for cell, width in zip(row, widths):
            if pdf.font_name_body == "Helvetica":
                cell = cell.replace("€", "EUR")
            align = "R" if cell[:1].isdigit() or cell[:1] == "-" else "L"
            pdf.cell(width, 7, cell, border=1, align=align)
        pdf.ln()


def payslip() -> Path:
    p = d.PAYSLIP
    pdf = _pdf()
    _text(pdf, "CEDOLINO PAGA - DOCUMENTO FITTIZIO PER DEMO", 14, bold=True, h=8)
    _text(pdf, f"Datore di lavoro: {p['employer']}")
    _text(pdf, f"Dipendente: {d.PERSON} - Matricola 000123 - Livello 4° CCNL Commercio")
    _text(pdf, f"Periodo di retribuzione: {p['month']}")
    pdf.ln(3)
    irpef_net = p["irpef_gross"] - p["irpef_deductions"]
    taxable = p["gross"] - p["inps"]
    _table(pdf, [
        ("Voce", "Competenze", "Trattenute"),
        ("Retribuzione lorda mensile (paga base + contingenza)", eur(p["gross"]), ""),
        ("Contributi previdenziali INPS c/dipendente (9,19%)", "", eur(p["inps"])),
        (f"Imponibile fiscale IRPEF: {eur(taxable)}", "", ""),
        ("IRPEF lorda", "", eur(p["irpef_gross"])),
        ("Detrazioni lavoro dipendente (art. 13 TUIR)", eur(p["irpef_deductions"]), ""),
        ("IRPEF netta trattenuta", "", eur(irpef_net)),
        ("Addizionale regionale IRPEF (rata 9/11)", "", eur(p["regional_surcharge"])),
        ("Addizionale comunale IRPEF (rata 9/11)", "", eur(p["municipal_surcharge"])),
    ], (120, 35, 35))
    pdf.ln(3)
    _text(pdf, f"NETTO IN BUSTA: {eur(p['net'])}", 13, bold=True, h=8)
    pdf.ln(2)
    _text(pdf, (
        f"Quota TFR maturata nel mese: {eur(p['tfr_month'])} (accantonata, non corrisposta in busta). "
        "Ratei tredicesima e quattordicesima maturati secondo CCNL. "
        "Conguaglio fiscale di fine anno a cura del sostituto d'imposta."
    ), 8, h=4)
    _text(pdf, "Documento generato con dati inventati: nessun riferimento a persone o aziende reali.", 8, h=4)
    return _save(pdf, "busta_paga_demo.pdf")


def bank_statement() -> Path:
    pdf = _pdf()
    _text(pdf, "ESTRATTO CONTO CORRENTE - DOCUMENTO FITTIZIO PER DEMO", 14, bold=True, h=8)
    _text(pdf, "Banca Esempio S.p.A. (banca fittizia) - Conto n. 0000 1234 5678")
    _text(pdf, f"Intestatario: {d.PERSON} - Periodo: 01/09/2026 - 30/09/2026")
    _text(pdf, "Dettaglio addebiti del periodo", 11, bold=True, h=8)
    rows = [("Data", "Descrizione operazione", "Addebito")]
    for day, item in enumerate(d.BANK_EXPENSES, start=1):
        label = item["label"]
        if item["periodicity"] == "trimestrale":
            label += " (addebito trimestrale)"
        rows.append((f"{min(day * 2, 30):02d}/09/2026", label, eur(item["amount"])))
    _table(pdf, rows, (28, 122, 40))
    pdf.ln(3)
    _text(pdf, "Note informative", 10, bold=True)
    _text(pdf, (
        "Finanziamento auto n. 778899: TAN 7,90% - TAEG 8,45%. Rate mensili con addebito diretto SDD. "
        "In caso di ritardato pagamento sono applicati interessi di mora come da contratto. "
        "Commissione di gestione: canone mensile di tenuta conto. "
        "Imposta di bollo: addebitata trimestralmente nella misura prevista dalla legge. "
        "Valuta: data da cui decorrono gli interessi sul movimento."
    ), 8, h=4)
    _text(pdf, "Documento generato con dati inventati: nessun riferimento a persone o banche reali.", 8, h=4)
    return _save(pdf, "estratto_conto_demo.pdf")


def expense_sheet() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Spese e entrate annuali"
    ws.append(["Voce", "Importo (EUR)", "Frequenza", "Tipo", "Note"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for item in d.SHEET_EXPENSES:
        ws.append([item["label"], item["amount"], item["periodicity"], "uscita", ""])
    extra = d.INCOMES[1]
    ws.append([extra["label"], extra["amount"], extra["periodicity"], "entrata", "compensi netti dopo la ritenuta del 20%"])
    for column, width in zip("ABCDE", (50, 15, 12, 10, 40)):
        ws.column_dimensions[column].width = width
    path = OUT_DIR / "foglio_spese_demo.xlsx"
    wb.save(path)
    return path


def _save(pdf: FPDF, name: str) -> Path:
    path = OUT_DIR / name
    pdf.output(str(path))
    return path


if __name__ == "__main__":
    OUT_DIR.mkdir(exist_ok=True)
    for path in (payslip(), bank_statement(), expense_sheet()):
        print(f"generato {path.relative_to(APP_DIR.parent)}")
